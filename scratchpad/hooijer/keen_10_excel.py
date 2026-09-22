"""Bouwt de KEEN-deepdive als Excel. Waarden, geen formules: LibreOffice kan in
deze omgeving niet doorrekenen en een formule zonder uitkomst leest als leeg."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import sql

L = "Arial"
KOP = Font(name=L, bold=True, color="FFFFFF", size=10)
KOPV = PatternFill("solid", fgColor="2F4858")
GEW, VET = Font(name=L, size=10), Font(name=L, bold=True, size=10)
LICHT = PatternFill("solid", fgColor="EDF2F4")
RAND = Border(bottom=Side(style="thin", color="BBBBBB"))
EURO, EURO2 = '€ #,##0;(€ #,##0);-', '€ #,##0.00;(€ #,##0.00);-'
PROC, AANT = '0.00"%";(0.00"%");-', '#,##0;(#,##0);-'
SCHOON = "SINCE 2026-03-01 UNTIL 2026-08-31"
M = "sessions, sessions_with_cart_additions, sessions_that_reached_checkout, sessions_that_completed_checkout"


def blad(wb, naam, titel, df, fmt, ondertitel=None):
    ws = wb.create_sheet(naam)
    ws["A1"], ws["A1"].font = titel, VET
    if ondertitel:
        ws["A2"], ws["A2"].font = ondertitel, GEW
    start = 4 if ondertitel else 3
    for k, n in enumerate(df.columns, 1):
        c = ws.cell(row=start, column=k, value=n)
        c.font, c.fill = KOP, KOPV
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
    for r, (_, reg) in enumerate(df.iterrows(), start + 1):
        laatste = str(reg.iloc[0]).startswith(("TOTAAL", "ALLE", "SAMEN"))
        for k, n in enumerate(df.columns, 1):
            c = ws.cell(row=r, column=k, value=reg[n])
            c.font, c.border = (VET if laatste else GEW), RAND
            if laatste:
                c.fill = LICHT
            if n in fmt:
                c.number_format = fmt[n]
    for k, n in enumerate(df.columns, 1):
        b = max(len(str(n)) + 2, 12)
        if k == 1:
            b = min(max(len(str(x)) for x in df.iloc[:, 0]) + 3, 50)
        ws.column_dimensions[get_column_letter(k)].width = b
    ws.row_dimensions[start].height = 30
    ws.freeze_panes = ws.cell(row=start + 1, column=2)
    return ws


def af(df, kol="sessions"):
    df = df.copy()
    df["wagen%"] = df.sessions_with_cart_additions / df[kol] * 100
    if "sessions_that_completed_checkout" in df:
        df["conv%"] = df.sessions_that_completed_checkout / df[kol] * 100
    return df.replace([float("inf"), float("-inf")], 0).fillna(0)


wb = Workbook()
ws = wb.active
ws.title = "Verantwoording"
for r, (t, f) in enumerate([
    ("KEEN NL — deepdive", VET), ("", GEW),
    ("Periode: 1 maart t/m 31 augustus 2026 (zes hele maanden, buiten de meetstoring).", GEW),
    ("Bron: Shopify Analytics via ShopifyQL, en de Shopify Admin API voor assortiment", GEW),
    ("en orders. Opgehaald op 22 september 2026.", GEW), ("", GEW),
    ("Brutomarge", VET),
    ("Gerekend met 54%. Shopify geeft voor KEEN 55,6% in juli en 54,2% in augustus;", GEW),
    ("dat zijn de enige maanden waarin de kostprijzen gevuld en stabiel zijn. Maart tot", GEW),
    ("en met mei staan op nul. Bij andere shops is de kostprijsdata nog onbruikbaar.", GEW),
    ("", GEW),
    ("Retourafhandeling", VET),
    ("Gerekend met € 9,00 per retourpakket (retourzending, inspectie, herverpakken).", GEW),
    ("Dit is een aanname, geen gemeten bedrag. Vervang hem door het echte cijfer.", GEW),
    ("", GEW),
    ("Wat getoetst is en niet klopte", VET),
    ("De aanname dat de lage conversie door uitverkochte maten komt, houdt geen stand.", GEW),
    ("Productpagina's met 76-100% van de maten op voorraad scoren niet beter dan", GEW),
    ("pagina's met 1-25%. Ook op alleen augustus gemeten, dicht bij de voorraadopname,", GEW),
    ("is er geen verband. HEYDUDE heeft 86% van de producten volledig uitverkocht en", GEW),
    ("converteert drie keer beter. Voorraad is dus niet de knop.", GEW),
], 1):
    ws.cell(row=r, column=1, value=t).font = f
ws.column_dimensions["A"].width = 92

# trechter per maand
m = af(sql.vraag("keen_nl", f"FROM sessions SHOW {M} GROUP BY month {SCHOON} ORDER BY month"))
m["Maand"] = m.month.str[:7]
m = m[["Maand", "sessions", "sessions_with_cart_additions", "wagen%",
       "sessions_that_reached_checkout", "sessions_that_completed_checkout", "conv%"]]
m.columns = ["Maand", "Sessies", "In wagen", "wagen%", "Afrekenen", "Order", "conv%"]
blad(wb, "Trechter", "KEEN NL: trechter per maand",
     m, {"Sessies": AANT, "In wagen": AANT, "Afrekenen": AANT, "Order": AANT,
         "wagen%": PROC, "conv%": PROC})

# verkeersbron, KEEN naast HEYDUDE
rijen = []
for s, n in (("keen_nl", "KEEN NL"), ("heydude_nl", "HEYDUDE NL")):
    d = af(sql.vraag(s, f"FROM sessions SHOW {M} GROUP BY referrer_source {SCHOON} "
                        f"ORDER BY sessions DESC"))
    d["Shop"] = n
    d["aandeel%"] = d.sessions / d.sessions.sum() * 100
    rijen.append(d)
b = pd.concat(rijen)[["Shop", "referrer_source", "sessions", "aandeel%", "wagen%", "conv%"]]
b.columns = ["Shop", "Bron", "Sessies", "aandeel%", "wagen%", "conv%"]
blad(wb, "Verkeersbron", "Verkeersbron: KEEN naast HEYDUDE",
     b, {"Sessies": AANT, "aandeel%": PROC, "wagen%": PROC, "conv%": PROC},
     "KEEN draait op zoekverkeer (55,7%), HEYDUDE op direct en social. Binnen zoekverkeer "
     "doet KEEN 3,27% tegen HEYDUDE 12,30%.")

# landingspagina's
p = af(sql.vraag("keen_nl", f"FROM sessions SHOW {M} GROUP BY landing_page_path {SCHOON} "
                            f"ORDER BY sessions DESC LIMIT 40"))
p = p[["landing_page_path", "sessions", "sessions_with_cart_additions", "wagen%", "conv%"]]
p.columns = ["Landingspagina", "Sessies", "In wagen", "wagen%", "conv%"]
blad(wb, "Paginas", "KEEN NL: 40 grootste landingspagina's",
     p, {"Sessies": AANT, "In wagen": AANT, "wagen%": PROC, "conv%": PROC},
     "Collectiepagina's zijn 27% van het verkeer en scoren 3,18%; bij HEYDUDE 9,05%.")

# retouren
r = sql.vraag("keen_nl", f"FROM sales SHOW gross_sales, returns, net_sales, net_items_sold "
                         f"GROUP BY product_title {SCHOON} ORDER BY returns ASC LIMIT 30")
r["retour%"] = -r.returns / r.gross_sales * 100
r = r[["product_title", "gross_sales", "returns", "retour%", "net_sales"]]
r.columns = ["Product", "Bruto-omzet", "Retouren", "retour%", "Netto-omzet"]
blad(wb, "Retouren", "KEEN NL: producten met de meeste retouren",
     r, {"Bruto-omzet": EURO, "Retouren": EURO, "Netto-omzet": EURO, "retour%": PROC},
     "Shopgemiddelde 22,2%. Retouren kosten KEEN € 85.306 brutowinst op jaarbasis.")

# ingrepen
ing = pd.DataFrame([
    ["1. Collectiepagina's van 3,18% naar 4,5%", "42.946 sessies, 3,18% in de wagen",
     256, 10165, 20330],
    ["2. Zes modellen boven 30% retour naar 22,2%", "€ 5.072 te veel retour op zes modellen",
     0, 3360, 6720],
    ["3. Terugkerende klanten naar retourniveau nieuw", "behoudt nu 59,3% tegen 77,8% bij nieuw",
     0, 8293, 16586],
    ["SAMEN", "", 256, 21818, 43636],
], columns=["Ingreep", "Nulmeting", "Extra orders per half jaar",
            "Brutowinst per half jaar", "Brutowinst per jaar"])
blad(wb, "Ingrepen", "KEEN NL: drie ingrepen, met de nulmeting om over een maand op terug te kijken",
     ing, {"Extra orders per half jaar": AANT, "Brutowinst per half jaar": EURO,
           "Brutowinst per jaar": EURO},
     "Ter vergelijking: de eerste trechterstap +1 procentpunt is € 57.258 brutowinst per jaar.")

wb.save("keen_deepdive.xlsx")
print("geschreven: keen_deepdive.xlsx —", wb.sheetnames)
