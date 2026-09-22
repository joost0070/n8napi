"""Werkdocument Bartogi, versie met seizoen: wat verkoopt er nog dit kwartaal?"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

L = "Arial"
KOP = Font(name=L, bold=True, color="FFFFFF", size=10); KOPV = PatternFill("solid", fgColor="2F4858")
GEW, VET = Font(name=L, size=10), Font(name=L, bold=True, size=10)
ROOD = PatternFill("solid", fgColor="FCE4E4"); GEEL = PatternFill("solid", fgColor="FFF6DA")
GROEN = PatternFill("solid", fgColor="E8F4EA"); BLAUW = PatternFill("solid", fgColor="E6EEF5")
RAND = Border(bottom=Side(style="thin", color="CCCCCC"))
EURO, EURO2 = '€ #,##0;(€ #,##0);-', '€ #,##0.00;(€ #,##0.00);-'
PROC, AANT, GETAL = '0.0"%";(0.0"%");-', '#,##0;(#,##0);-', '0.0'

d = pd.read_pickle("bar_q4_realistisch.pkl").copy()

KOL = [("shop","Shop",None),("merk","Merk",None),("artikel","Artikel",None),
       ("soort","Productsoort",None),("seizoensindex","Seizoens-index",GETAL),
       ("voorraad","Voorraad",AANT),("stuks","Verkocht jan-sep",AANT),
       ("verwacht okt-dec","Verwacht okt-dec",GETAL),
       ("verkoopbaar Q4","Verkoopbaar Q4",GETAL),
       ("retour%","Retour%",PROC),("prijs nu incl","Prijs nu incl",EURO2),
       ("vanprijs incl","Adviesprijs incl",EURO2),("afprijzing%","Afprijzing%",PROC),
       ("kostprijs","Inkoopprijs",EURO2),("bruto","Omzet jan-sep",EURO),
       ("marge%","Marge%",PROC),("bodemprijs incl","Bodemprijs 20% incl",EURO2),
       ("doelprijs","Doelprijs (max adviesprijs)",EURO2),
       ("verhoging","Verhoging",EURO2),
       ("kans realistisch","Kans Q4 in euro",EURO),
       ("advies_seizoen","Advies",None)]


def tabel(ws, df, start=4):
    for k, (_, naam, _) in enumerate(KOL, 1):
        c = ws.cell(row=start, column=k, value=naam)
        c.font, c.fill = KOP, KOPV
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
    for i, (_, r) in enumerate(df.iterrows(), start + 1):
        for k, (veld, _, fmt) in enumerate(KOL, 1):
            v = r[veld]
            c = ws.cell(row=i, column=k, value=(None if pd.isna(v) else v))
            c.font, c.border = GEW, RAND
            if fmt: c.number_format = fmt
            if veld == "marge%" and pd.notna(v):
                c.fill = GROEN if v >= 20 else (GEEL if v >= 0 else ROOD)
            if veld == "kans realistisch" and pd.notna(v) and v > 0:
                c.fill = BLAUW
    br = {"Artikel": 48, "Advies": 52, "Merk": 12, "Shop": 6, "Productsoort": 20}
    for k, (_, naam, _) in enumerate(KOL, 1):
        ws.column_dimensions[get_column_letter(k)].width = br.get(naam, 13)
    ws.row_dimensions[start].height = 34
    ws.freeze_panes = ws.cell(row=start + 1, column=4)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(KOL))}{start+len(df)}"


wb = Workbook(); ws = wb.active; ws.title = "Verantwoording"
for i, (t, f) in enumerate([
    ("Bartogi — marge, seizoen en bodemprijs per artikel", VET), ("", GEW),
    ("Periode: verkopen 1 januari t/m 21 september 2026. Bedragen exclusief btw", GEW),
    ("tenzij een kolom 'incl' heet. Voorraad is de stand van 22 september 2026.", GEW), ("", GEW),
    ("Seizoen", VET),
    ("Het jaargemiddelde is voor schoenen een slechte voorspeller. Daarom is per", GEW),
    ("productsoort uit 2025 afgeleid hoeveel er in oktober t/m december verkocht", GEW),
    ("werd per stuk dat in januari t/m september ging:", GEW),
    ("     Pantoffels 1,40   Laarzen/boots 0,89   Sneakers 0,23   Instappers 0,19", GEW),
    ("     Wandelschoenen 0,17   Slippers 0,07   Sandalen 0,06   Espadrilles 0,01", GEW),
    ("Verwacht okt-dec = verkocht jan-sep x die index. Verkoopbaar Q4 is dat getal,", GEW),
    ("begrensd door de voorraad.", GEW), ("", GEW),
    ("Bodemprijs en doelprijs", VET),
    ("Bodemprijs = de prijs waarbij dit artikel 20% brutowinstmarge haalt, met zijn", GEW),
    ("eigen inkoopprijs en retourpercentage:", GEW),
    ("   P = ( inkoop x (1-retour%) + verzendkosten/stuk - verzendopbrengst/stuk )", GEW),
    ("       / ( (1-retour%) - payment% - advertentie% - 20% )", GEW),
    ("Doelprijs = de bodemprijs, maar nooit hoger dan de eigen adviesprijs. Waar de", GEW),
    ("bodemprijs boven de adviesprijs ligt is het met prijs niet op te lossen.", GEW),
    ("Kans Q4 = verkoopbaar Q4 x (doelprijs - huidige prijs), exclusief btw.", GEW), ("", GEW),
    ("Advies", VET),
    ("  op doel                                    marge 20% of hoger", GEW),
    ("  verkoopt dit kwartaal uit: prijs omhoog    vraag dit kwartaal haalt de voorraad in", GEW),
    ("  prijs omhoog, gefaseerd                    voorraad 1,5 tot 4 kwartalen vraag", GEW),
    ("  veel meer voorraad dan Q4-vraag            meer dan 4 kwartalen vraag op de plank;", GEW),
    ("        prijs handhaven en ruimen", GEW),
    ("  buiten seizoen                             minder dan 1 stuk vraag dit kwartaal;", GEW),
    ("        de afprijzing levert nu niets op en verbrandt de prijs voor het voorjaar", GEW),
    ("  bodemprijs boven adviesprijs               niet met prijs op te lossen; inkoop-", GEW),
    ("        voorwaarden of assortimentskeuze heroverwegen", GEW),
    ("  onhaalbaar bij huidige kosten              retour zo hoog dat geen prijs 20% haalt", GEW),
    ("", GEW), ("Let op", VET),
    ("De retourzending zelf (ongeveer € 6 per pakket) zit hier niet in, net zomin als", GEW),
    ("in de MT-rapportage. De werkelijke marge ligt dus iets lager.", GEW),
], 1):
    ws.cell(row=i, column=1, value=t).font = f
ws.column_dimensions["A"].width = 96

nl = d[d.shop == "NL"]
ws = wb.create_sheet("Vooruitzicht Q4")
ws["A1"] = "Wat er in oktober t/m december nog te verdienen valt"; ws["A1"].font = VET
regels = [("", "Bartogi NL", "Bartogi DE"),
          ("voorraad (stuks)", 45689, 30290),
          ("verwachte vraag okt-dec (stuks)", 1039, 984),
          ("daarvan op voorraad", 704, 576),
          ("Q4-omzet bij huidige prijs", 40496, 40267),
          ("Q4-brutowinst bij ongewijzigd beleid", 5484, 14067),
          ("prijsruimte binnen de adviesprijs", 9034, 2865),
          ("advertentie naar 12% in Q4", 3612, 0),
          ("", "", ""),
          ("jaar 2026 bij ongewijzigd beleid", 0.0820, 0.2451),
          ("jaar 2026 met beide ingrepen", 0.1374, 0.2561)]
for i, (a, b, c) in enumerate(regels, 3):
    ws.cell(row=i, column=1, value=a).font = VET if i in (3, 12, 13) else GEW
    for k, v in ((2, b), (3, c)):
        if v == "": continue
        cel = ws.cell(row=i, column=k, value=v); cel.font = GEW
        cel.number_format = '0.00"%"' if isinstance(v, float) and v < 1 else AANT
        if isinstance(v, (int, float)) and v > 1000: cel.number_format = EURO
        if i in (4, 5, 6): cel.number_format = AANT
for k, b in (("A", 40), ("B", 16), ("C", 16)):
    ws.column_dimensions[k].width = b
ws.cell(row=12, column=2).number_format = '0.00%'
ws.cell(row=12, column=3).number_format = '0.00%'
ws.cell(row=13, column=2).number_format = '0.00%'
ws.cell(row=13, column=3).number_format = '0.00%'

ws = wb.create_sheet("Kans dit kwartaal")
ws["A1"] = "Artikelen met prijsruimte die dit kwartaal nog verkopen (Bartogi NL en DE)"
ws["A1"].font = VET
ws["A2"] = "Gesorteerd op wat de prijsverhoging dit kwartaal oplevert."; ws["A2"].font = GEW
tabel(ws, d[d["kans realistisch"] > 0].sort_values("kans realistisch", ascending=False).head(120))

ws = wb.create_sheet("Buiten seizoen")
ws["A1"] = "Afgeprijsd, maar verkoopt dit kwartaal toch niet"; ws["A1"].font = VET
ws["A2"] = ("Deze artikelen staan gemiddeld 40% afgeprijsd terwijl er dit kwartaal "
            "nauwelijks vraag naar is. De korting levert nu niets op.")
ws["A2"].font = GEW
tabel(ws, d[d.advies_seizoen.str.startswith("buiten seizoen")].sort_values(
    "voorraad", ascending=False))

GROOT = ["KEEN", "Lazamani", "HEYDUDE", "Hunter", "Crocs", "Toni Pons", "Tofvel"]
d["merkgroep"] = d.merk.str.upper().map(
    lambda m: next((g for g in GROOT if g.upper() == m), "Overige merken"))
for groep in GROOT + ["Overige merken"]:
    sub = d[d.merkgroep == groep].sort_values("kans realistisch", ascending=False)
    if not len(sub): continue
    ws = wb.create_sheet(groep[:31])
    ws["A1"] = f"{groep}: marge, seizoen en bodemprijs per artikel"; ws["A1"].font = VET
    ws["A2"] = (f"{len(sub)} artikelen | omzet jan-sep € {sub.bruto.sum():,.0f} | "
                f"kans dit kwartaal € {sub['kans realistisch'].sum():,.0f}").replace(",", ".")
    ws["A2"].font = GEW
    tabel(ws, sub)

wb.save("bartogi_marge_per_artikel.xlsx")
print("geschreven:", wb.sheetnames)
