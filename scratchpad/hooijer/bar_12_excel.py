"""Werkdocument Bartogi: marge en bodemprijs per artikel, per merk."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

L = "Arial"
KOP = Font(name=L, bold=True, color="FFFFFF", size=10)
KOPV = PatternFill("solid", fgColor="2F4858")
GEW, VET = Font(name=L, size=10), Font(name=L, bold=True, size=10)
ROOD = PatternFill("solid", fgColor="FCE4E4")
GEEL = PatternFill("solid", fgColor="FFF6DA")
GROEN = PatternFill("solid", fgColor="E8F4EA")
RAND = Border(bottom=Side(style="thin", color="CCCCCC"))
EURO, EURO2 = '€ #,##0;(€ #,##0);-', '€ #,##0.00;(€ #,##0.00);-'
PROC, AANT, GETAL = '0.0"%";(0.0"%");-', '#,##0;(#,##0);-', '0.0'

d = pd.read_pickle("bar_werkdocument.pkl")
mod = pd.read_pickle("bar_modellen_nl.pkl")

KOLOMMEN = [("shop","Shop",None),("merk","Merk",None),("artikel","Artikel",None),
            ("maten","Maten",AANT),("voorraad","Voorraad",AANT),
            ("voorraaddekking mnd","Voorraad in mnd",GETAL),
            ("stuks","Verkocht",AANT),("retour%","Retour%",PROC),
            ("prijs nu incl","Prijs nu incl",EURO2),("vanprijs incl","Adviesprijs incl",EURO2),
            ("afprijzing%","Afprijzing%",PROC),("kostprijs","Inkoopprijs",EURO2),
            ("bruto","Omzet excl",EURO),("brutowinst","Brutowinst",EURO),
            ("marge%","Marge%",PROC),("tekort","Tekort tot 20%",EURO),
            ("bodemprijs incl","Bodemprijs 20% incl",EURO2),
            ("verhoging nodig","Verhoging nodig",EURO2),("advies","Advies",None)]


def tabel(ws, df, start=3):
    for k, (_, naam, _) in enumerate(KOLOMMEN, 1):
        c = ws.cell(row=start, column=k, value=naam)
        c.font, c.fill = KOP, KOPV
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
    for i, (_, r) in enumerate(df.iterrows(), start + 1):
        for k, (veld, _, fmt) in enumerate(KOLOMMEN, 1):
            v = r[veld]
            c = ws.cell(row=i, column=k, value=(None if pd.isna(v) else v))
            c.font, c.border = GEW, RAND
            if fmt:
                c.number_format = fmt
            if veld == "marge%" and pd.notna(v):
                c.fill = GROEN if v >= 20 else (GEEL if v >= 0 else ROOD)
    breedtes = {"Artikel": 48, "Advies": 46, "Merk": 12, "Shop": 6}
    for k, (_, naam, _) in enumerate(KOLOMMEN, 1):
        ws.column_dimensions[get_column_letter(k)].width = breedtes.get(naam, 13)
    ws.row_dimensions[start].height = 30
    ws.freeze_panes = ws.cell(row=start + 1, column=4)
    ws.auto_filter.ref = (f"A{start}:{get_column_letter(len(KOLOMMEN))}"
                          f"{start + len(df)}")


wb = Workbook()
ws = wb.active; ws.title = "Verantwoording"
for i, (t, f) in enumerate([
    ("Bartogi — marge en bodemprijs per artikel", VET), ("", GEW),
    ("Periode: 1 januari t/m 21 september 2026. Bedragen exclusief btw tenzij anders vermeld.", GEW),
    ("Bron: alle orderregels uit Shopify (bartogi.nl en bartogi.de) met de kostprijs per", GEW),
    ("variant, gekoppeld aan de adviesprijs (compareAtPrice) en de voorraad. Kostenposten", GEW),
    ("komen uit de MT-rapportage 'Omzet en kosten Week en YTD'.", GEW), ("", GEW),
    ("Hoe de marge per artikel is berekend", VET),
    ("  omzet (na afprijzing en korting)", GEW),
    ("  min de omzet die als retour terugkwam", GEW),
    ("  min de inkoopprijs van wat verkocht bleef", GEW),
    ("  plus verzendopbrengsten, min verzendkosten (€ 4,85 per stuk NL, € 4,88 DE)", GEW),
    ("  min payment costs (0,78% NL, 3,10% DE van de omzet)", GEW),
    ("  min advertentiekosten (20,92% NL, 8,78% DE van de omzet)", GEW),
    ("  = brutowinst. Marge% = brutowinst gedeeld door de bruto omzet, zoals in de MT-rapportage.", GEW),
    ("", GEW),
    ("Bodemprijs", VET),
    ("De verkoopprijs waarbij dit artikel exact 20% brutowinstmarge haalt, met zijn eigen", GEW),
    ("inkoopprijs en zijn eigen retourpercentage:", GEW),
    ("   P = ( inkoop x (1-retour%) + verzendkosten/stuk - verzendopbrengst/stuk )", GEW),
    ("       / ( (1-retour%) - payment% - advertentie% - 20% )", GEW),
    ("Bij een advertentiedruk van 20,92% komt dat neer op ongeveer twee keer de inkoopprijs.", GEW),
    ("Zakt de advertentiedruk naar 12%, dan is 1,65 keer de inkoopprijs al genoeg.", GEW),
    ("", GEW),
    ("Advies-kolom", VET),
    ("  op doel                          marge is 20% of hoger", GEW),
    ("  prijs omhoog naar bodemprijs     onder doel, voorraad is binnen 6 maanden weg", GEW),
    ("  prijs omhoog, gefaseerd          onder doel, voorraad 6 tot 18 maanden", GEW),
    ("  voorraad te groot: prijs handhaven, snel weg   meer dan 18 maanden voorraad; de", GEW),
    ("        opbrengst van doorverkopen weegt zwaarder dan de marge per stuk", GEW),
    ("  bodemprijs ligt boven de adviesprijs           met prijs alleen niet op te lossen;", GEW),
    ("        inkoopvoorwaarden of assortimentskeuze heroverwegen", GEW),
    ("  onhaalbaar bij huidige kosten: uitfaseren      retourpercentage zo hoog dat geen", GEW),
    ("        prijs 20% haalt", GEW),
    ("", GEW),
    ("Let op", VET),
    ("De retourzending zelf (ongeveer € 6 per retourpakket) zit niet in deze berekening,", GEW),
    ("net zomin als in de MT-rapportage. De werkelijke marge ligt dus iets lager.", GEW),
], 1):
    ws.cell(row=i, column=1, value=t).font = f
ws.column_dimensions["A"].width = 100

# samenvatting
ws = wb.create_sheet("Samenvatting")
nl = d[d.shop == "NL"]
rijen = [
    ("Waar de marge van Bartogi NL is gebleven", "", ""),
    ("", "", ""),
    ("Post", "Bedrag", "% van bruto omzet"),
    ("afprijzing (verkoopprijs onder adviesprijs)", 102011, 54.5),
    ("advertentiekosten", 39188, 20.9),
    ("retouren (omzet die terugkwam)", 35938, 19.2),
    ("verzendkosten minus verzendopbrengsten", 14205, 7.6),
    ("", "", ""),
    ("Verdeling van de artikelen", "aantal", "% van de omzet"),
]
for i, (a, b, c) in enumerate(rijen, 1):
    ws.cell(row=i, column=1, value=a).font = VET if i in (1, 3, 9) else GEW
    if b != "":
        cel = ws.cell(row=i, column=2, value=b); cel.font = GEW
        if isinstance(b, (int, float)): cel.number_format = EURO
    if c != "":
        cel = ws.cell(row=i, column=3, value=c); cel.font = GEW
        if isinstance(c, (int, float)): cel.number_format = PROC
g = nl.groupby("advies").agg(artikelen=("artikel","count"), omzet=("bruto","sum"),
                             tekort=("tekort","sum")).reset_index()
g["aandeel"] = g.omzet / nl.bruto.sum() * 100
for j, (_, r) in enumerate(g.sort_values("tekort", ascending=False).iterrows(), 10):
    ws.cell(row=j, column=1, value=r.advies).font = GEW
    ws.cell(row=j, column=2, value=int(r.artikelen)).font = GEW
    c = ws.cell(row=j, column=3, value=float(r.aandeel)); c.font = GEW; c.number_format = PROC
    c = ws.cell(row=j, column=4, value=float(r.tekort)); c.font = GEW; c.number_format = EURO
ws.cell(row=9, column=4, value="tekort tot 20%").font = VET
for k, b in (("A", 52), ("B", 14), ("C", 16), ("D", 16)):
    ws.column_dimensions[k].width = b

# top drukkers
ws = wb.create_sheet("Top drukkers")
ws["A1"] = "De 40 artikelen die de brutowinstmarge het hardst drukken (Bartogi NL)"
ws["A1"].font = VET
ws["A2"] = "Tekort = wat er ontbrak om op 20% brutowinstmarge uit te komen."
ws["A2"].font = GEW
tabel(ws, nl.sort_values("tekort", ascending=False).head(40), start=4)

# retouren
ws = wb.create_sheet("Retouren")
ws["A1"] = "Artikelen met de grootste retourschade (Bartogi NL, minimaal 10 stuks verkocht)"
ws["A1"].font = VET
ws["A2"] = "Gesorteerd op de omzet die als retour terugkwam."
ws["A2"].font = GEW
ret = nl[nl.stuks >= 10].sort_values("retour", ascending=False).head(40)
tabel(ws, ret, start=4)

# per merk
GROOT = ["KEEN", "Lazamani", "HEYDUDE", "Hunter", "Crocs", "Toni Pons", "Tofvel"]
d["merkgroep"] = d.merk.str.upper().map(lambda m: next(
    (g for g in GROOT if g.upper() == m), "Overige merken"))
for groep in GROOT + ["Overige merken"]:
    sub = d[d.merkgroep == groep].sort_values("tekort", ascending=False)
    if not len(sub):
        continue
    ws = wb.create_sheet(groep[:31])
    ws["A1"] = f"{groep}: marge en bodemprijs per artikel"
    ws["A1"].font = VET
    ws["A2"] = (f"{len(sub)} artikelen | omzet € {sub.bruto.sum():,.0f} | "
                f"brutowinst € {sub.brutowinst.sum():,.0f} | "
                f"marge {sub.brutowinst.sum()/sub.bruto.sum()*100:.1f}%").replace(",", ".")
    ws["A2"].font = GEW
    tabel(ws, sub, start=4)

wb.save("bartogi_marge_per_artikel.xlsx")
print("geschreven:", wb.sheetnames)
print(f"totaal {len(d):,} artikelregels".replace(",", "."))
