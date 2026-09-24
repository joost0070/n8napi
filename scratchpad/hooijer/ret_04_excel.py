"""Bouwt de retouranalyse met losse redenen als Excel, per shopgroep."""
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from ret_03_analyse import analyse, tabellen, KOLOM

L = "Arial"
KOP = Font(name=L, bold=True, color="FFFFFF", size=10); KOPV = PatternFill("solid", fgColor="2F4858")
GEW, VET, TITEL = Font(name=L, size=10), Font(name=L, bold=True, size=10), Font(name=L, bold=True, size=12)
ROOD = PatternFill("solid", fgColor="FCE4E4"); GEEL = PatternFill("solid", fgColor="FFF6DA")
GROEN = PatternFill("solid", fgColor="E8F4EA"); RAND = Border(bottom=Side(style="thin", color="CCCCCC"))
PROC, AANT, EURO = '0.0', '#,##0', '€ #,##0'


def schrijf(ws, df, start, formaten, breed=None, kleur_retour=None):
    for k, n in enumerate(df.columns, 1):
        c = ws.cell(row=start, column=k, value=str(n)); c.font, c.fill = KOP, KOPV
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, (_, r) in enumerate(df.iterrows(), start + 1):
        for k, n in enumerate(df.columns, 1):
            v = r[n]
            if isinstance(v, pd.Timestamp): v = v.date()
            c = ws.cell(row=i, column=k, value=None if (not isinstance(v, str) and pd.isna(v)) else v)
            c.font, c.border = GEW, RAND
            if n in formaten: c.number_format = formaten[n]
            if kleur_retour and n == "Retour%" and not pd.isna(v):
                c.fill = ROOD if v >= kleur_retour + 10 else (GROEN if v < kleur_retour - 10 else GEEL)
            if isinstance(n, str) and n.startswith("% ") and not pd.isna(v) and v >= 40:
                c.font = VET
    for k, n in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(k)].width = (breed or {}).get(n, max(10, min(16, len(str(n)) + 3)))
    ws.row_dimensions[start].height = 32
    ws.freeze_panes = ws.cell(row=start + 1, column=2)
    return start + len(df)


def bouw(cfg, pad):
    v, ret, los = analyse(cfg)
    t = tabellen(v, ret, cfg)
    wb = Workbook()

    # samenvatting
    ws = wb.active; ws.title = "Samenvatting"
    ws["A1"] = f"{cfg['naam']}, retouranalyse met losse redenen"; ws["A1"].font = TITEL
    m, rd = t["model"], t["redenen"]
    landen = sorted(v.land.unique())
    regels = [("KERNCIJFERS", None),
              ("Meetvenster", f"bestellingen {cfg['van']} t/m {cfg['tot']}, retouren daarop t/m {cfg['retouren_tot']}"),
              ("Verkochte paren", int(t["verkocht"])), ("Bestellingen", int(v.order.nunique())),
              ("Geretourneerde paren", int(len(ret))), ("Retourpercentage", f"{t['gem']:.1f}%")]
    for land in landen:
        rl = (ret.land == land).sum() / v[v.land == land].aantal.sum() * 100
        regels.append((f"Retourpercentage {land}", f"{rl:.1f}%"))
    regels += [("Omzet in het venster", round(v.netto.sum())), ("Retourwaarde", round(ret.prijs.sum())),
               ("Retouren die een ruil werden", f"{ret.ruil.mean()*100:.0f}%"), ("", None),
               ("REDENEN, ELK APART (% van alle retouren)", None)]
    for _, r in rd.iterrows():
        regels.append((f"  {r.Reden}", f"{r['% van retouren']:.1f}%   ({int(r.Totaal)} paren, ruil {r['Ruil%']:.0f}%)"))
    regels += [("", None), ("WAT DE REDENEN ZEGGEN, per model met minstens 10 retouren", None)]
    diag = m[m.Retour >= 10].groupby("Wat de redenen zeggen").agg(
        modellen=("model", lambda s: ", ".join(s.head(8))), retouren=("Retour", "sum")).sort_values(
        "retouren", ascending=False)
    for d, r in diag.iterrows():
        regels.append((f"  {d}", f"{int(r.retouren)} retouren: {r.modellen}"))
    regels += [("", None), ("HOE TE LEZEN", None),
               ("Retour%", "geretourneerde paren uit het venster / verkochte paren uit het venster"),
               ("% te klein, % te smal, ...", "aandeel van die reden in de retouren van dat model"),
               ("Prioriteit Aanpakken", f"retour% minstens 10 punt boven het gemiddelde, of een van de drie modellen met de hoogste retourwaarde en boven het gemiddelde; minstens 30 verkocht"),
               ("Waarom los", "een maat groter lost 'te klein' op, niet 'te smal' of 'instap te nauw'"),
               ("Benchmark", "Returnista Europa per categorie: sandalen 23,2%, boots/laarzen 25,1%, winter-/wandelschoenen 28,5%, sneakers 17,5-30%, sokken 1,25%, footwear totaal 28%. Teenslippers en espadrilles vergeleken met sandalen; soorten zonder eigen benchmark met footwear totaal. De benchmark rekent mogelijk per stuk of per order; hier per paar."),
               ("Niet gekoppeld", f"{los} retourregels vielen buiten het venster of misten een sku")]
    for i, (a, b) in enumerate(regels, 3):
        ws.cell(row=i, column=1, value=a).font = VET if (b is None and a) else GEW
        if b is not None:
            c = ws.cell(row=i, column=2, value=b); c.font = GEW
            if isinstance(b, int) and b > 10000: c.number_format = EURO
            elif isinstance(b, int): c.number_format = AANT
    ws.column_dimensions["A"].width = 62; ws.column_dimensions["B"].width = 110

    # per model
    ws = wb.create_sheet("Retour% per model")
    ws["A1"] = "Echt retourpercentage per model, met elke reden apart"; ws["A1"].font = TITEL
    ws["A2"] = f"Gemiddeld {t['gem']:.1f}%. Redenkolommen = aandeel van die reden in de retouren van het model. Vet = 40% of meer."
    ws["A2"].font = GEW
    kol = (["model", "Soort", "Verkocht", "Retour", "Retour%", "Boven gemiddelde",
            "Benchmark", "T.o.v. benchmark", "Benchmark van"]
           + [c for l in landen for c in (f"Verkocht {l}", f"Retour% {l}")]
           + list(KOLOM.values()) + ["% overig", "Ruil%", "Omzet", "Retourwaarde",
                                    "Wat de redenen zeggen", "Thema's in toelichting", "Prioriteit"])
    df = m[kol].rename(columns={"model": "Model"})
    f = {c: PROC for c in df.columns if "%" in c or c in ("Boven gemiddelde", "Benchmark", "T.o.v. benchmark")}
    f.update({"Verkocht": AANT, "Retour": AANT, "Omzet": EURO, "Retourwaarde": EURO})
    f.update({f"Verkocht {l}": AANT for l in landen})
    schrijf(ws, df, 4, f, breed={"Model": 16, "Soort": 18, "Benchmark van": 20, "Wat de redenen zeggen": 58,
                                 "Thema's in toelichting": 52, "Prioriteit": 28}, kleur_retour=t["gem"])

    # per maat
    ws = wb.create_sheet("Per maat")
    ws["A1"] = "Retourpercentage per maat, met elke reden apart"; ws["A1"].font = TITEL
    s = t["maat"].rename(columns={"maat": "Maat"})
    eind = schrijf(ws, s, 3, {c: PROC for c in s.columns if "%" in c} | {"Verkocht": AANT, "Retour": AANT})
    r0 = eind + 3
    ws.cell(row=r0, column=1, value="Retourpercentage per model en maat (leeg = minder dan 5 verkocht)").font = VET
    p = t["pct"].reset_index().rename(columns={"model": "Model"})
    p.columns = ["Model"] + [f"Maat {c}" for c in p.columns[1:]]
    eind = schrijf(ws, p, r0 + 1, {c: '0' for c in p.columns[1:]})
    r1 = eind + 3
    ws.cell(row=r1, column=1, value="Verkochte paren per model en maat (de onderbouwing van de tabel hierboven)").font = VET
    n = t["n"].reset_index().rename(columns={"model": "Model"})
    n.columns = ["Model"] + [f"n {c}" for c in n.columns[1:]]
    schrijf(ws, n, r1 + 1, {c: AANT for c in n.columns[1:]})

    # redenen
    ws = wb.create_sheet("Redenen")
    ws["A1"] = f"Retourredenen, {' tegenover '.join(landen)}"; ws["A1"].font = TITEL
    rd = t["redenen"]
    schrijf(ws, rd, 3, {"% van retouren": PROC, "% van verkoop": PROC, "Retourwaarde": EURO,
                        "Ruil%": PROC} | {l: AANT for l in landen} | {"Totaal": AANT},
            breed={"Reden": 32})

    # model x reden
    ws = wb.create_sheet("Model x reden")
    ws["A1"] = "Aantal geretourneerde paren per model en reden"; ws["A1"].font = TITEL
    mx = t["mxr"]
    schrijf(ws, mx, 3, {c: AANT for c in mx.columns[1:]}, breed={"Model": 16})

    # per week
    ws = wb.create_sheet("Retour% per week")
    ws["A1"] = "Retourpercentage per bestelweek"; ws["A1"].font = TITEL
    wk = t["week"]
    schrijf(ws, wk, 3, {c: PROC for c in wk.columns if "%" in c} | {c: AANT for c in wk.columns
                                                                     if c.startswith(("Verkocht", "Retour ")) and "%" not in c})

    wb.save(pad)
    return t, ret, v, los


LAZ = r"^(\S+)"
TOF = r"^(Mula Lungta|Mula|Rabara Maha|Rabara|Slipa Maha|Slipa|Luna Kids|Luna|Sundara|Celsi|Kaya|Bodhi)"
CONFIG = {
    "lazamani": dict(naam="Lazamani NL + DE", shops=[("lazamani_nl", "NL"), ("lazamani_de", "DE")],
                     van="2026-06-01", tot="2026-08-19", retouren_tot="2026-08-31", model=LAZ,
                     retouren=lambda: pd.read_pickle("retouren.pkl").query("shopcode in ['LANL','LADE']")),
    "tofvel": dict(naam="Tofvel NL + DE", shops=[("tofvel_nl", "NL"), ("tofvel_de", "DE")],
                   van="2026-01-01", tot="2026-09-23", retouren_tot="2026-09-23", model=TOF,
                   retouren=lambda: pd.read_pickle("tofvel_retouren.pkl")),
}

if __name__ == "__main__":
    for naam in sys.argv[1:]:
        t, ret, v, los = bouw(CONFIG[naam], f"retouranalyse_{naam}_redenen_los.xlsx")
        print(f"\n=== {CONFIG[naam]['naam']}: {int(t['verkocht']):,} paren verkocht, "
              f"{len(ret):,} retour, {t['gem']:.1f}% (niet gekoppeld: {los}) ===".replace(",", "."))
        print(t["redenen"][["Reden", "Totaal", "% van retouren", "Ruil%"]].to_string(
            index=False, float_format=lambda x: f"{x:,.1f}"))
