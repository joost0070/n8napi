"""Tofvel-deepdive: onderbouwende tabellen naar tofvel_deepdive_2026-09.xlsx (waarden)."""
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

NAAM = {"tofvel_nl": "Tofvel NL", "tofvel_de": "Tofvel DE", "tofvel_en": "Tofvel EU"}
PER = {"P": "29 jun-26 sep 2026", "V": "31 mrt-28 jun 2026", "J": "29 jun-26 sep 2025"}


def schoon(df):
    df = df.copy()
    if "shop" in df:
        df["shop"] = df.shop.map(NAAM).fillna(df.shop)
    if "periode" in df:
        df["periode"] = df.periode.map(PER).fillna(df.periode)
    return df.round(2)


k = schoon(pd.read_pickle("tdd_kern.pkl"))
t = pd.read_pickle("tdd_trechter.pkl")
tr = schoon(t[t.dim.isin(["totaal", "apparaat", "kanaal"])][
    ["shop", "periode", "dim", "waarde", "sessions", "sessions_with_cart_additions",
     "sessions_that_reached_checkout", "sessions_that_completed_checkout", "wagen_%",
     "checkout_%_van_wagen", "order_%_van_checkout", "conversie_%", "bot_%"]])
wk = schoon(t[t.dim == "week"][["shop", "periode", "waarde", "sessions", "wagen_%",
                                "checkout_%_van_wagen", "order_%_van_checkout", "conversie_%"]])
ow = pd.read_pickle("tdd_orderwinst.pkl")
grp = []
for dim in ["groep", "kanaal"]:
    g = ow.groupby(["shop", "periode", dim]).agg(
        orders=("order", "count"), omzet_ex=("netto_ex", "sum"), aov_ex=("netto_ex", "mean"),
        stuks=("stuks", "sum"), stuks_retour=("stuks_retour", "sum"),
        geld_terug_ex=("terug_ex", "sum"), winst=("winst", "sum"),
        winst_per_order=("winst", "mean")).reset_index().rename(columns={dim: "waarde"})
    g.insert(2, "dim", dim)
    grp.append(g)
grp = pd.concat(grp)
grp["retour_stuks_%"] = 100 * grp.stuks_retour / grp.stuks
grp = schoon(grp)
kl = schoon(pd.read_pickle("tdd_klanten.pkl"))
pr = pd.read_pickle("tdd_producten.pkl").sort_values("winst", ascending=False).round(2)
pk = pd.read_pickle("tdd_piekrisico.pkl")
pk = pk[["titel", "maat", "sku", "prijs", "kostprijs", "voorraad", "vraag", "tekort",
         "ruimte_andere_kleur", "winst_per_paar"]].rename(
    columns={"vraag": "verkocht_okt_dec_2025"}).sort_values(
    ["tekort", "verkocht_okt_dec_2025"], ascending=False).round(2)
prijs = schoon(pd.read_pickle("tdd_prijs.pkl"))
sh = pd.read_pickle("tdd_shops.pkl").drop(columns=["sleutel", "kost_dekking_%"]).round(2)
mt = pd.read_pickle("tdd_mt_ytd.pkl").round(3)
co = schoon(pd.read_pickle("tdd_checkouts.pkl"))
ret = pd.read_pickle("tofvel_retouren.pkl")
ret = ret[ret["Created at"].astype(str).str[:10] >= "2026-06-29"][
    ["shopcode", "Order number", "Product name", "Return reasons description",
     "return_reason_comment", "Resolution type"]]

VERANTWOORDING = [
    ("Periodes", "P = 29 jun-26 sep 2026 (90 dagen t/m gisteren); V = 31 mrt-28 jun 2026; "
                 "J = 29 jun-26 sep 2025. Seizoen: Tofvel verkoopt ~75% van de jaaromzet in okt-jan; "
                 "alle drie de vensters zijn laagseizoen, V is voorjaar."),
    ("Bedragen", "Excl. btw, per order omgerekend (btw/totaal). Alleen webshoporders; omruilorders "
                 "van Returnista (bedrag ~€0) apart geteld."),
    ("Kostprijs", "Huidige unitCost per variant (100% gevuld). Shopify legt kostprijs pas sinds 2026 "
                  "vast bij verkoop; voor 2024/25 en 2025 is de huidige kostprijs gebruikt (aanname: "
                  "inkoopprijs niet wezenlijk veranderd)."),
    ("Retouren", "Retourbedrag = werkelijk terugbetaald geld (excl. terugbetaalde verzendkosten), "
                 "cohort: orders uit de periode. Retourstuks tellen omruilingen mee. Retourverzendkosten "
                 "niet bekend en niet meegenomen."),
    ("Kosten", "Payment % van bruto en verzendkosten per verzonden order uit MT-rapport (NL 0,70% en "
               "€6,00; DE 5,04% en €6,50). Advertentie = Google/Bing/Meta uit KPI-rapport; week 39 2026 "
               "ontbreekt nog en is geschat als 6/7 van week 38."),
    ("Sessies", "Botfilter: alleen afzetlanden (NL: Nederland + België; DE: Duitsland, Oostenrijk, "
                "Zwitserland, Luxemburg; EU: Europese afzetlanden excl. Ierland). In P is 85% van alle "
                "NL-sessies nep (Brazilië, Pakistan, Zuid-Afrika e.a., 0 winkelwagens). Vergelijking met "
                "andere shops: filter 'landen met minstens één order' voor alle shops gelijk."),
    ("Meetdekking", "Shopify koppelt in P maar 48% van de NL-orders aan een sessie (J: 69%); 43% van de "
                    "orders heeft geen herkomst (geen cookietoestemming). Conversie op sessies is dus "
                    "conversie van bezoekers die toestemming gaven."),
    ("Productpagina", "ShopifyQL levert geen stap 'productpagina bekeken'; etalage-analyse gebruikt "
                      "binnenkomst op productpagina als benadering."),
    ("Piekrisico", "Vraag = verkochte paren okt-dec 2025 (NL+DE+EU, gedeelde voorraad). Tekort = vraag - "
                   "huidige voorraad per SKU. Bijdrage per paar = prijs excl. btw - kostprijs - €6 verzending "
                   "- payment. Bovengrens: gaat uit van gelijke vraag per kleur en maat, zonder uitwijk."),
    ("Bronnen", "Shopify Admin API (orders, producten, voorraad, verlaten checkouts), ShopifyQL (sessies, "
                "verkoop, zoekopdrachten), Returnista-export t/m 23 sep 2026, MT-rapport week 38, Tofvel "
                "KPI-rapport t/m week 38."),
]

TABS = [("Kerncijfers", k), ("Trechter", tr), ("Trechter per week", wk),
        ("Nieuw-terug & kanaal", grp), ("Klantwaarde cohorts", kl),
        ("Producten P", pr), ("Piekrisico per maat", pk), ("Prijs per maand", prijs),
        ("Retouren P (Returnista)", ret), ("Verlaten checkouts P", co),
        ("Alle shops P", sh), ("MT YTD alle kanalen", mt)]

with pd.ExcelWriter("tofvel_deepdive_2026-09.xlsx", engine="openpyxl") as x:
    pd.DataFrame(VERANTWOORDING, columns=["Onderwerp", "Toelichting"]).to_excel(
        x, sheet_name="Verantwoording", index=False)
    for naam, df in TABS:
        df.to_excel(x, sheet_name=naam, index=False)
    for ws in x.book.worksheets:
        for rij in ws.iter_rows():
            for c in rij:
                c.font = Font(name="Arial", size=10, bold=c.row == 1,
                              color="FFFFFF" if c.row == 1 else "000000")
                if c.row == 1:
                    c.fill = PatternFill("solid", fgColor="1F4E78")
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, kol in enumerate(ws.columns, 1):
            breedte = max(len(str(c.value or "")) for c in list(kol)[:200])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, breedte + 2), 60)
        ws.freeze_panes = "A2"
    vw = x.book["Verantwoording"]
    vw.column_dimensions["B"].width = 120
    for c in vw["B"]:
        c.alignment = Alignment(wrap_text=True, vertical="top")
print("klaar")
